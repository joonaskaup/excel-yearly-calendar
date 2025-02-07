import pandas as pd
import calendar
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_calendar(input_file, output_file, config):
    # Extract settings from config
    title = config.get("GENERAL", {}).get("TITLE", None)
    column_width = float(config.get("GENERAL", {}).get("COLUMN_WIDTH", 4.5))
    phase_colors = config.get("PHASE_COLORS", {})
    row_heights = config.get("ROW_HEIGHTS", {"normal": "20", "special": "50"})

    # Normalize phase keys and convert row heights to floats
    phase_colors = {k.strip().lower(): v for k, v in phase_colors.items()}
    row_heights = {k.lower(): float(v) for k, v in row_heights.items()}

    # Read the input Excel file
    df = pd.read_excel(input_file)

    # Convert 'Start' and 'End' to datetime
    df['Start'] = pd.to_datetime(df['Start'], dayfirst=True, errors='coerce')
    df['End'] = pd.to_datetime(df['End'], dayfirst=True, errors='coerce')
    df = df.dropna(subset=['Start'])
    df['End'] = df['End'].fillna(df['Start'])
    df['Phase'] = df['Phase'].astype(str).str.strip().str.lower()

    years = sorted(set(df['Start'].dt.year).union(df['End'].dt.year))

    wb = Workbook()
    ws = wb.active
    ws.title = "Yearly Calendar"

    # Define styles
    title_font = Font(bold=True, size=24)
    year_header_font = Font(bold=True, color='FFFFFF', size=20)
    year_header_fill = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
    month_header_font = Font(bold=True, color='000000')
    month_header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    week_header_font = Font(bold=True, color='000000')
    week_header_fill = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
    event_font = Font(color='000000')
    event_fill = PatternFill(fill_type=None)
    alignment_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    alignment_left = Alignment(horizontal='left', vertical='center', wrap_text=False)
    thin_border = Border(
        left=Side(border_style="thin", color="000000"),
        right=Side(border_style="thin", color="000000"),
        top=Side(border_style="thin", color="000000"),
        bottom=Side(border_style="thin", color="000000")
    )

    # Set column widths for 48 columns
    for col in range(1, 49):
        ws.column_dimensions[get_column_letter(col)].width = column_width

    current_row = 1

    if title:
        title_cell = ws.cell(row=current_row, column=1, value=title)
        title_cell.font = title_font
        title_cell.alignment = alignment_center
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=48)
        ws.row_dimensions[current_row].height = 30
        current_row += 1

    for year in years:
        months = [(year, m) for m in range(1, 13)]
        month_positions = {(year, m): idx for idx, (year, m) in enumerate(months)}

        def date_to_week_index(date):
            if date.year != year:
                return None
            month_key = (date.year, date.month)
            if month_key not in month_positions:
                return None
            month_index = month_positions[month_key]
            day = date.day
            if day <= 7:
                week_in_month = 0
            elif day <= 14:
                week_in_month = 1
            elif day <= 21:
                week_in_month = 2
            else:
                week_in_month = 3
            return month_index * 4 + week_in_month

        df_year = df[(df['Start'].dt.year <= year) & (df['End'].dt.year >= year)].copy()
        df_year['Start_week'] = df_year['Start'].apply(date_to_week_index).fillna(0).astype(int)
        df_year['End_week'] = df_year['End'].apply(date_to_week_index).fillna(47).astype(int)
        df_sorted = df_year.sort_values(by='Start_week')

        long_events = []
        short_events = []
        for _, event in df_sorted.iterrows():
            duration = event['End_week'] - event['Start_week'] + 1
            if duration <= 2:
                short_events.append(event)
            else:
                long_events.append(event)

        # Assign long events to rows avoiding overlaps
        rows_assigned = []
        for event in long_events:
            assigned = False
            for row_list in rows_assigned:
                if all(event['End_week'] < e['Start_week'] or event['Start_week'] > e['End_week'] for e in row_list):
                    row_list.append(event)
                    assigned = True
                    break
            if not assigned:
                rows_assigned.append([event])

        # Year header row
        year_cell = ws.cell(row=current_row, column=1, value=str(year))
        year_cell.font = year_header_font
        year_cell.fill = year_header_fill
        year_cell.alignment = alignment_center
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=48)
        for col in range(1, 49):
            ws.cell(row=current_row, column=col).border = thin_border
        ws.row_dimensions[current_row].height = 30
        current_row += 1

        # Month header row
        col = 1
        for ym in months:
            month_name = calendar.month_name[ym[1]]
            month_cell = ws.cell(row=current_row, column=col, value=month_name)
            month_cell.font = month_header_font
            month_cell.fill = month_header_fill
            month_cell.alignment = alignment_center
            ws.merge_cells(start_row=current_row, start_column=col, end_row=current_row, end_column=col+3)
            for c in range(col, col+4):
                ws.cell(row=current_row, column=c).border = thin_border
            col += 4
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        # Week header row
        col = 1
        for _ in months:
            for week in range(1, 5):
                week_cell = ws.cell(row=current_row, column=col, value=f"W{week}")
                week_cell.font = week_header_font
                week_cell.fill = week_header_fill
                week_cell.alignment = alignment_center
                week_cell.border = thin_border
                col += 1
        ws.row_dimensions[current_row].height = 15
        current_row += 1

        # Long events rows
        normal_height = row_heights.get("normal", 20)
        for events_row in rows_assigned:
            for col in range(1, 49):
                ws.cell(row=current_row, column=col).border = thin_border
                ws.cell(row=current_row, column=col).alignment = alignment_center
            for event in events_row:
                start_col = event['Start_week'] + 1
                end_col = event['End_week'] + 1
                if start_col > end_col:
                    start_col, end_col = end_col, start_col
                start_col = max(1, min(start_col, 48))
                end_col = max(1, min(end_col, 48))
                ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
                cell = ws.cell(row=current_row, column=start_col)
                cell.value = event['Title']
                cell.font = event_font
                cell.alignment = alignment_center
                fill_color = phase_colors.get(event['Phase'], None)
                if fill_color:
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
                else:
                    cell.fill = event_fill
                for idx in range(start_col, end_col + 1):
                    ws.cell(row=current_row, column=idx).border = thin_border
                    ws.cell(row=current_row, column=idx).alignment = alignment_center
            ws.row_dimensions[current_row].height = normal_height
            current_row += 1

        # Special row for short events
        if short_events:
            for col in range(1, 49):
                ws.cell(row=current_row, column=col).border = thin_border
                ws.cell(row=current_row, column=col).alignment = alignment_center
            for event in short_events:
                start_col = event['Start_week'] + 1
                end_col = event['End_week'] + 1
                if start_col > end_col:
                    start_col, end_col = end_col, start_col
                start_col = max(1, min(start_col, 48))
                end_col = max(1, min(end_col, 48))
                ws.merge_cells(start_row=current_row, start_column=start_col, end_row=current_row, end_column=end_col)
                cell = ws.cell(row=current_row, column=start_col)
                cell.value = event['Title']
                cell.font = event_font
                cell.alignment = alignment_center
                fill_color = phase_colors.get(event['Phase'], None)
                if fill_color:
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
                else:
                    cell.fill = event_fill
                for idx in range(start_col, end_col + 1):
                    ws.cell(row=current_row, column=idx).border = thin_border
                    ws.cell(row=current_row, column=idx).alignment = alignment_center
            special_height = row_heights.get("special", 50)
            ws.row_dimensions[current_row].height = special_height
            current_row += 1

        # Blank row between years
        current_row += 1

    # Legend at bottom
    legend_start = current_row + 1
    legend_col = 1
    legend_title = ws.cell(row=legend_start, column=legend_col, value="Legend:")
    legend_title.font = Font(bold=True)
    ws.row_dimensions[legend_start].height = 20
    legend_row = legend_start + 1
    for phase, color in phase_colors.items():
        color_cell = ws.cell(row=legend_row, column=legend_col)
        color_cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        color_cell.border = thin_border
        color_cell.alignment = alignment_center
        ws.column_dimensions[get_column_letter(legend_col)].width = column_width
        phase_cell = ws.cell(row=legend_row, column=legend_col + 1, value=phase.title())
        phase_cell.alignment = alignment_left
        phase_cell.border = thin_border
        ws.merge_cells(start_row=legend_row, start_column=legend_col + 1, end_row=legend_row, end_column=legend_col + 5)
        legend_row += 1

    wb.save(output_file)